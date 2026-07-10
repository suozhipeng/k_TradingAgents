"""Tests for AStock PPTX report generation (Phase 17).

Verifies:
- ReportGenerator initialization (requires python-pptx)
- to_pptx() produces a valid .pptx file
- Slide count and structure (cover, research, advisory, provider, trace)
- Content correctness (symbol, trade_date, status)
- Graceful handling when python-pptx is absent
- API endpoint for PPTX download
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

# ---------------------------------------------------------------------------
# Sample report data
# ---------------------------------------------------------------------------

SAMPLE_REPORT = {
    "symbol": "600519.SH",
    "normalized_symbol": "600519.SH",
    "trade_date": "2026-06-15",
    "source": "test",
    "status": "ok",
    "summary": "Sample research report for testing.",
    "investment_plan": "Buy on dips below 1800.",
    "runtime_profile": "test_profile",
    "runtime_trace": [
        "Step 1: Initialise runtime",
        "Step 2: Fetch market data",
        "Step 3: Run analysis",
        "Step 4: Generate report",
    ],
    "provider_coverage": {
        "eastmoney": {"status": "ok", "source": "eastmoney", "records": 120},
        "sina": {"status": "ok", "source": "sina", "records": 85},
        "cninfo": {"status": "error", "source": "cninfo", "records": 0},
    },
    "missing_data_notes": ["Research provider returned empty data."],
    "degradation_notes": [],
    "research_conclusion": {
        "recommendation": "buy",
        "confidence": "high",
        "summary": "Strong fundamentals support upside.",
    },
    "trader_proposal": {
        "action": "buy",
        "quantity": 100,
        "price_limit": 1850.0,
        "rationale": "Momentum bullish.",
    },
    "risk_decision": {
        "approved": True,
        "max_position_pct": 5.0,
        "stop_loss_pct": 7.0,
        "notes": "Within risk limits.",
    },
    "portfolio_decision": {
        "action": "buy",
        "quantity": 80,
        "target_weight_pct": 3.0,
        "portfolio_notes": "Adding to existing position.",
    },
    "metadata": {},
}


# ---------------------------------------------------------------------------
# Tests: PPT report generation
# ---------------------------------------------------------------------------


class TestReportGenerator:
    """Tests for ReportGenerator class."""

    def test_import_check(self):
        """HAS_PPTX should be boolean."""
        from tradingagents.astock.reporting.ppt import HAS_PPTX

        assert isinstance(HAS_PPTX, bool)

    def test_raises_without_pptx(self):
        """ReportGenerator should raise ImportError when python-pptx not available."""
        with patch("tradingagents.astock.reporting.ppt.HAS_PPTX", False):
            with pytest.raises(ImportError, match="python-pptx"):
                from tradingagents.astock.reporting.ppt import ReportGenerator

                ReportGenerator()

    @pytest.mark.skipif(
        not __import__("importlib").util.find_spec("pptx"),
        reason="python-pptx not installed",
    )
    def test_to_pptx_creates_file(self):
        """to_pptx() should create a valid .pptx file."""
        from tradingagents.astock.reporting.ppt import ReportGenerator

        gen = ReportGenerator()
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            result = gen.to_pptx(SAMPLE_REPORT, tmp_path)
            assert result == tmp_path
            assert os.path.exists(tmp_path)
            assert os.path.getsize(tmp_path) > 1000  # non-trivial file

            # Verify it's a valid ZIP (PPTX is a ZIP archive)
            import zipfile
            with zipfile.ZipFile(tmp_path, "r") as zf:
                names = zf.namelist()
                assert any("slide" in n for n in names), "No slides in PPTX"
        finally:
            Path(tmp_path).unlink(missing_ok=True)

    @pytest.mark.skipif(
        not __import__("importlib").util.find_spec("pptx"),
        reason="python-pptx not installed",
    )
    def test_slide_count(self):
        """Should produce exactly 5 slides."""
        from tradingagents.astock.reporting.ppt import ReportGenerator
        from pptx import Presentation

        gen = ReportGenerator()
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            gen.to_pptx(SAMPLE_REPORT, tmp_path)
            prs = Presentation(tmp_path)
            assert len(prs.slides) == 5, f"Expected 5 slides, got {len(prs.slides)}"
        finally:
            Path(tmp_path).unlink(missing_ok=True)

    @pytest.mark.skipif(
        not __import__("importlib").util.find_spec("pptx"),
        reason="python-pptx not installed",
    )
    def test_slides_contain_expected_text(self):
        """Slides should contain key text from the report."""
        from tradingagents.astock.reporting.ppt import ReportGenerator
        from pptx import Presentation

        gen = ReportGenerator()
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            gen.to_pptx(SAMPLE_REPORT, tmp_path)
            prs = Presentation(tmp_path)

            # Collect all text from all slides
            all_text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        all_text += shape.text + "\n"

            assert "600519.SH" in all_text, "Symbol missing from slides"
            assert "2026-06-15" in all_text, "Trade date missing from slides"
            assert "buy" in all_text.lower(), "Recommendation missing from slides"
            assert "Research Conclusion" in all_text
            assert "Provider Coverage" in all_text
            assert "Runtime Trace" in all_text
        finally:
            Path(tmp_path).unlink(missing_ok=True)


class TestReportsAPI:
    """Tests for the /api/v1/reports/pptx endpoint."""

    @pytest.fixture
    def app(self):
        from tradingagents.astock.api import create_app

        app = create_app(db_path=":memory:", cors_origin="*", test_config={"ASTOCK_RESEARCH_ONLY": False})
        app.config["TESTING"] = True
        return app

    @pytest.fixture
    def client(self, app):
        return app.test_client()

    def test_pptx_endpoint_returns_501_without_pptx(self, client):
        """Should return 501 if python-pptx is not importable."""
        with patch("tradingagents.astock.api.routes_reports.HAS_PPTX", False):
            resp = client.get("/api/v1/reports/pptx?symbol=600519.SH")
            assert resp.status_code == 501
            data = resp.get_json()
            assert data is not None
            assert "python-pptx" in data.get("error", "") or "python-pptx" in str(data)

    def test_pptx_endpoint_returns_pptx(self, client):
        """Should return a PPTX file when python-pptx is available."""
        resp = client.get("/api/v1/reports/pptx?symbol=600519.SH")
        if resp.status_code == 501:
            pytest.skip("python-pptx not installed on this system")
        assert resp.status_code == 200
        assert resp.mimetype == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        assert len(resp.data) > 1000
